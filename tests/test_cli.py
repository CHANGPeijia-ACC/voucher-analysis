from pathlib import Path

import pytest
from openpyxl import load_workbook

from voucher_analysis import cli
from voucher_analysis import reconcile as rec

REPO_ROOT = Path(__file__).resolve().parents[1]
CONFIG = str(REPO_ROOT / "config.yaml")
SAMPLE = REPO_ROOT / "data" / "sample"


def run_arguments(out_dir, ground_truth=None):
    arguments = ["run", "--gl", str(SAMPLE / "gl.csv"), "--bank", str(SAMPLE / "bank.csv"),
                 "--config", CONFIG, "--out", str(out_dir)]
    if ground_truth:
        arguments += ["--ground-truth", str(SAMPLE / "ground_truth.csv")]
    return arguments


def test_help_lists_both_commands(capsys):
    with pytest.raises(SystemExit):
        cli.main(["--help"])
    text = capsys.readouterr().out

    assert "generate" in text
    assert "run" in text


def test_command_is_required(capsys):
    with pytest.raises(SystemExit):
        cli.main([])


def test_generate_writes_the_four_files(tmp_path, capsys):
    code = cli.main(["generate", "--out", str(tmp_path), "--seed", "7"])
    written = sorted(path.name for path in tmp_path.glob("*.csv"))

    assert code == 0
    assert written == ["bank.csv", "bank_ground_truth.csv", "gl.csv", "ground_truth.csv"]
    assert "vouchers" in capsys.readouterr().out


def test_run_writes_a_workbook_with_evaluation(tmp_path, capsys):
    code = cli.main(run_arguments(tmp_path, ground_truth=True))
    workbook = load_workbook(tmp_path / cli.WORKBOOK_NAME)
    printed = capsys.readouterr().out

    assert code == 0
    assert "Evaluation" in workbook.sheetnames
    assert "JET01" in printed
    assert "precision" in printed


def test_run_without_ground_truth_skips_evaluation(tmp_path):
    code = cli.main(run_arguments(tmp_path))
    workbook = load_workbook(tmp_path / cli.WORKBOOK_NAME)

    assert code == 0
    assert "Evaluation" not in workbook.sheetnames
    assert "BankRec" in workbook.sheetnames


def test_run_creates_the_output_folder(tmp_path):
    out_dir = tmp_path / "new" / "folder"
    assert cli.main(run_arguments(out_dir)) == 0
    assert (out_dir / cli.WORKBOOK_NAME).exists()


def test_run_returns_one_when_the_statement_does_not_tie(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(rec, "statement_ties", lambda statement: False)
    code = cli.main(run_arguments(tmp_path))

    assert code == 1
    assert "does not tie" in capsys.readouterr().out
