from pathlib import Path

import pytest
from openpyxl import load_workbook

from gl_builder import make_bank, make_gl, voucher
from voucher_analysis import jet, reconcile as rec, report
from voucher_analysis.loaders import load_config

REPO_ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def config():
    return load_config(REPO_ROOT / "config.yaml")


@pytest.fixture
def workbook(tmp_path, config):
    """A small workbook built the same way the command line will build it."""
    gl = make_gl(
        *voucher("JV000001", 60000, description="Office rent"),
        {"voucher_no": "JV000002", "line_no": 1, "debit": 200.00},
        {"voucher_no": "JV000002", "line_no": 2, "debit": None, "credit": 190.00},
        *voucher("JV000003", 1000, debit_account="2202", credit_account="1002",
                 posting_date="2025-12-30"),
    )
    bank = make_bank({"bank_date": "2025-12-31", "amount": -150,
                      "counterparty": "Bank service charge"})
    flags = jet.run_all(gl, config)
    result = rec.reconcile(gl, bank, config)
    path = report.write_workbook(tmp_path / "workpaper.xlsx", gl, flags,
                                 jet.month_end_summary(gl, config), result)
    return load_workbook(path)


def test_summary_counts_each_voucher_once(config):
    gl = make_gl(*voucher("JV000001", 60000))          # 60,000 debit and 60,000 credit
    flags = jet.run_all(gl, config)
    summary = report.summary_table(flags, gl).set_index("test_id")

    assert summary.loc["JET06", "lines_flagged"] == 2
    assert summary.loc["JET06", "vouchers_flagged"] == 1
    assert summary.loc["JET06", "voucher_debit_flagged"] == 60000.00   # not 120,000
    assert summary.loc["JET01", "lines_flagged"] == 0


def test_summary_amount_covers_tests_that_flag_only_a_credit_line(config):
    """JET08 flags the bank line of a payment, which has no debit."""
    rows = []
    for day, number in [("03", 1), ("04", 2), ("05", 3)]:
        rows += voucher(f"JV00000{number}", 45000, debit_account="2202",
                        credit_account="1002", posting_date=f"2025-03-{day}")
    gl = make_gl(*rows)
    summary = report.summary_table(jet.run_all(gl, config), gl).set_index("test_id")

    assert summary.loc["JET08", "lines_flagged"] == 3
    assert summary.loc["JET08", "voucher_debit_flagged"] == 135000.00


def test_detail_has_ledger_fields_and_reason(config):
    gl = make_gl(*voucher("JV000001", 60000, description="Office rent"))
    detail = report.test_detail(jet.run_all(gl, config), gl, "JET06")

    assert list(detail.columns) == report.DETAIL_COLUMNS
    assert detail["description"].iloc[0] == "Office rent"
    assert detail["reason"].iloc[0].startswith("Round amount")


def test_unmatched_table_includes_amount_differences(config):
    gl = make_gl(*voucher("JV000001", 10000, debit_account="2202", credit_account="1002",
                          posting_date="2025-03-03"))
    bank = make_bank({"bank_date": "2025-03-04", "amount": -10025,
                      "counterparty": "Supplier A"})
    table = report.unmatched_table(rec.reconcile(gl, bank, config))

    assert list(table.columns) == report.UNMATCHED_SHEET_COLUMNS
    assert list(table["item_type"]) == ["amount_difference"]
    assert table["difference"].iloc[0] == -25.00


def test_workbook_has_every_sheet(workbook):
    expected = ["Summary"] + list(report.TEST_DESCRIPTIONS) + ["MonthEnd", "BankRec", "Unmatched"]
    assert workbook.sheetnames == expected


def test_workbook_says_the_data_is_synthetic(workbook):
    assert workbook["Summary"]["A1"].value == report.SYNTHETIC_NOTE
    assert workbook["Summary"]["A3"].value == "test_id"       # header sits below the note


def test_workbook_formatting(workbook):
    summary = workbook["Summary"]
    assert summary["A3"].font.bold
    assert summary.freeze_panes == "A4"
    assert summary["E4"].number_format == report.MONEY_FORMAT

    detail = workbook["JET01"]
    assert detail.freeze_panes == "A2"
    assert detail["C2"].number_format == report.DATE_FORMAT   # posting_date
    assert detail["D2"].number_format == report.TIME_FORMAT   # entry_time


def test_workbook_keeps_headers_on_an_empty_sheet(workbook):
    sheet = workbook["JET00"]                                  # nothing to flag here
    assert sheet["A1"].value == "voucher_no"
    assert sheet.max_row == 1


def test_workbook_adds_evaluation_only_when_given(tmp_path, config):
    gl = make_gl(*voucher("JV000001", 60000))
    flags = jet.run_all(gl, config)
    result = rec.reconcile(gl, make_bank({"bank_date": "2025-03-03", "amount": -5}), config)
    evaluation = flags.head(0).assign(precision=[], recall=[])

    path = report.write_workbook(tmp_path / "with_evaluation.xlsx", gl, flags,
                                 jet.month_end_summary(gl, config), result, evaluation)
    assert "Evaluation" in load_workbook(path).sheetnames
